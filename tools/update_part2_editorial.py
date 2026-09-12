from pathlib import Path
from docx import Document
import json, hashlib

ROOT=Path(__file__).resolve().parents[1]
PKG=ROOT/'papers/2026_V1I1_PIONEER_JULY_AUGUST_PART2'
E035_NOTE=('Author contributions: The supplied manuscript establishes the three-author order but does not include a CRediT statement. '
           'Contribution roles must be confirmed by all authors in the final editorial proof; no unverified roles are assigned in this copy.')
E046_NOTE=('Author contributions: Chia, Shiaondo Kenneth, and Amenger, Gbaden: conceptualization, methodology, formal analysis, '
           'engineering reconstruction, validation, visualization, investigation, writing—original draft, and writing—review and editing. '
           'Goshwe, N. Y.: supervisory contribution and critical review as recorded in the supplied manuscript. All authors must confirm this statement in the final editorial proof.')

def insert_before(d, marker, text):
    if any(text.split(':',1)[0].lower() in p.text.lower() for p in d.paragraphs): return
    target=next((p for p in d.paragraphs if p.text.strip().lower()==marker.lower()),None)
    if target is None: d.add_paragraph(text); return
    new=target.insert_paragraph_before(text); new.style=target.style

def edit_doc(eid, note):
    doc=next((PKG/eid/'manuscript').glob('*.docx'))
    d=Document(doc); insert_before(d,'References',note); d.save(doc); return doc

def sha(p):
    h=hashlib.sha256(); h.update(p.read_bytes()); return h.hexdigest().upper()

def update_meta(eid, contributions, extra=None):
    path=PKG/eid/'metadata/article_metadata.json'; data=json.loads(path.read_text(encoding='utf-8'))
    data['author_contributions']=contributions
    if extra: data.update(extra)
    doc=next((PKG/eid/'manuscript').glob('*.docx')); data['manuscript_sha256']=sha(doc)
    path.write_text(json.dumps(data,indent=2,ensure_ascii=False),encoding='utf-8')
    with (PKG/eid/'editorial/PART2_RECONCILIATION.md').open('a',encoding='utf-8') as f:
        f.write('\n- Author contributions: '+contributions+'\n')

def main():
    edit_doc('e035',E035_NOTE); update_meta('e035',E035_NOTE,{'author_contribution_status':'AUTHOR_CONFIRMATION_REQUIRED'})
    edit_doc('e046',E046_NOTE); update_meta('e046',E046_NOTE,{'author_contribution_status':'EDITORIAL_STATEMENT_REQUIRES_ALL_AUTHOR_CONFIRMATION'})
    with (PKG/'e031/editorial/PART2_RECONCILIATION.md').open('a',encoding='utf-8') as f:
        f.write('\n- Manuscript control: the current e031 manuscript is a complete replacement of the legacy grouped copy. The current source hash is authoritative; the legacy copy is retained only for provenance and is not a publication candidate.\n')
    with (PKG/'e056/editorial/PART2_RECONCILIATION.md').open('a',encoding='utf-8') as f:
        f.write('\n- Editorial hold: the requested renewable-paper replacement source was not present in the controlled workspace. The supplied e056 RESTORE manuscript remains authoritative; do not relabel or substitute e044 without the intended renewable source and author authorization.\n')
    update_meta('e031',json.loads((PKG/'e031/metadata/article_metadata.json').read_text(encoding='utf-8')).get('author_contributions',''),{'manuscript_control':'CURRENT_MANUSCRIPT_REPLACES_LEGACY_GROUPED_COPY'})
    update_meta('e056',json.loads((PKG/'e056/metadata/article_metadata.json').read_text(encoding='utf-8')).get('author_contributions',''),{'editorial_hold':'RENEWABLE_SOURCE_NOT_SUPPLIED; supplied RESTORE manuscript preserved'})
    print('Updated e035, e046, e031 and e056 editorial controls.')

if __name__=='__main__': main()
